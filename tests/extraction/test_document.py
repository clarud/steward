from datetime import UTC, datetime
from pathlib import Path

from subprocess import CompletedProcess

from steward.extraction import DocxExtractor, HtmlExtractor, ImageOcrExtractor, PdfExtractor, PlainTextExtractor
from steward.sources import Source, SourceType


def test_plain_text_extractor_keeps_entire_file_provenance(tmp_path: Path) -> None:
    path = tmp_path / "note.txt"
    path.write_text("A TLB caches translations.", encoding="utf-8")
    source = Source(1, path, "a" * 64, SourceType.PLAIN_TEXT, path.stat().st_size,
                    datetime(2026, 9, 8, tzinfo=UTC), datetime(2026, 9, 8, tzinfo=UTC), datetime(2026, 9, 8, tzinfo=UTC))

    result = PlainTextExtractor().extract(source)

    assert result.fragments[0].location == "entire file"
    assert result.fragments[0].text == "A TLB caches translations."


def test_pdf_extractor_keeps_page_provenance(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "note.pdf"
    path.write_bytes(b"%PDF")
    source = Source(1, path, "a" * 64, SourceType.PDF, path.stat().st_size,
                    datetime(2026, 9, 8, tzinfo=UTC), datetime(2026, 9, 8, tzinfo=UTC), datetime(2026, 9, 8, tzinfo=UTC))

    class Page:
        def __init__(self, text): self._text = text
        def extract_text(self): return self._text
    class Reader:
        pages = [Page("first"), Page(""), Page("third")]
    monkeypatch.setattr("steward.extraction.document.PdfReader", lambda path: Reader())

    result = PdfExtractor().extract(source)

    assert [(fragment.text, fragment.location) for fragment in result.fragments] == [
        ("first", "page 1"), ("third", "page 3")
    ]


def test_docx_extractor_retains_paragraph_and_heading_provenance(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "notes.docx"
    document = Document()
    document.add_heading("Address translation", level=1)
    document.add_paragraph("A TLB caches translations.")
    document.add_paragraph("Page tables map virtual memory.")
    document.save(path)
    source = Source(
        1,
        path,
        "a" * 64,
        SourceType.DOCX,
        path.stat().st_size,
        datetime(2026, 9, 8, tzinfo=UTC),
        datetime(2026, 9, 8, tzinfo=UTC),
        datetime(2026, 9, 8, tzinfo=UTC),
    )

    result = DocxExtractor().extract(source)

    assert [(fragment.heading, fragment.text, fragment.location) for fragment in result.fragments] == [
        ("Address translation", "A TLB caches translations.", "paragraph 2"),
        ("Address translation", "Page tables map virtual memory.", "paragraph 3"),
    ]


def test_html_extractor_ignores_scripts_and_uses_nearest_heading(tmp_path: Path) -> None:
    path = tmp_path / "notes.html"
    path.write_text(
        "<h1>Virtual memory</h1><p>Address translation maps virtual pages.</p>"
        "<script>ignore_this()</script><h2>TLB</h2><p>Caches translations.</p>",
        encoding="utf-8",
    )
    source = Source(1, path, "a" * 64, SourceType.HTML, path.stat().st_size,
                    datetime(2026, 9, 8, tzinfo=UTC), datetime(2026, 9, 8, tzinfo=UTC), datetime(2026, 9, 8, tzinfo=UTC))

    result = HtmlExtractor().extract(source)

    assert [(fragment.heading, fragment.text, fragment.location) for fragment in result.fragments] == [
        ("Virtual memory", "Address translation maps virtual pages.", "section 1"),
        ("TLB", "Caches translations.", "section 2"),
    ]


def test_image_ocr_extractor_uses_local_tesseract_and_keeps_image_provenance(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "receipt.png"
    path.write_bytes(b"png")
    source = Source(1, path, "a" * 64, SourceType.IMAGE, path.stat().st_size,
                    datetime(2026, 9, 8, tzinfo=UTC), datetime(2026, 9, 8, tzinfo=UTC), datetime(2026, 9, 8, tzinfo=UTC))
    captured: dict[str, object] = {}

    def fake_run(arguments, **kwargs):
        captured["arguments"] = arguments
        return CompletedProcess(arguments, 0, stdout="Booking reference: ABC123\n", stderr="")

    monkeypatch.setattr("steward.extraction.document.run", fake_run)

    result = ImageOcrExtractor().extract(source)

    assert captured["arguments"] == ["tesseract", str(path), "stdout"]
    assert [(fragment.text, fragment.location) for fragment in result.fragments] == [
        ("Booking reference: ABC123", "image OCR")
    ]
