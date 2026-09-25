from pathlib import Path
from zipfile import ZipFile

from steward.extraction import ExtractionResult, MarkdownExtractor, SourceFragment, SourceFragmentRepository
from steward.sources import SourceType
from steward.sources import SourceRepository
from steward.sources.service import SourceService
from steward.storage import initialize_database


def test_source_service_refreshes_fragments_when_a_source_changes(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"
    initialize_database(database_path)
    vault = tmp_path / "vault"
    vault.mkdir()
    source_path = vault / "note.md"
    source_path.write_text("# First\nOriginal text.", encoding="utf-8")
    source_repository = SourceRepository(database_path)
    fragment_repository = SourceFragmentRepository(database_path)
    service = SourceService(
        source_repository=source_repository,
        fragment_repository=fragment_repository,
        markdown_extractor=MarkdownExtractor(),
    )
    service.scan_markdown_root(vault)
    source = source_repository.get_by_path(source_path.resolve())
    assert source is not None
    source_path.write_text("# Second\nUpdated text.", encoding="utf-8")

    service.scan_markdown_root(vault)

    fragments = fragment_repository.list_for_source(source.id or 0)
    assert [fragment.heading for fragment in fragments] == ["Second"]
    assert [fragment.text for fragment in fragments] == ["# Second\nUpdated text."]


def test_source_service_scans_and_extracts_plain_text_files(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "vault"; vault.mkdir()
    source_path = vault / "todo.txt"; source_path.write_text("Study address translation", encoding="utf-8")
    sources = SourceRepository(database_path)
    fragments = SourceFragmentRepository(database_path)
    service = SourceService(sources, fragments, MarkdownExtractor())

    result = service.scan_source_root(vault)
    source = sources.get_by_path(source_path.resolve())

    assert result.new == 1
    assert source is not None
    assert [fragment.text for fragment in fragments.list_for_source(source.id or 0)] == ["Study address translation"]


def test_source_service_scans_and_extracts_source_code_with_path_provenance(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "project"; vault.mkdir()
    source_path = vault / "scheduler.py"
    source_path.write_text("def schedule():\n    return 'ready'\n", encoding="utf-8")
    sources = SourceRepository(database_path)
    fragments = SourceFragmentRepository(database_path)

    SourceService(sources, fragments, MarkdownExtractor()).scan_source_root(vault)

    source = sources.get_by_path(source_path.resolve())
    assert source is not None and source.source_type is SourceType.CODE
    extracted = fragments.list_for_source(source.id or 0)
    assert extracted[0].location == "lines 1-2"
    assert "def schedule" in extracted[0].text


def test_source_service_extracts_pptx_slide_text_with_slide_provenance(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "slides"; vault.mkdir()
    source_path = vault / "lecture.pptx"
    with ZipFile(source_path, "w") as archive:
        archive.writestr(
            "ppt/slides/slide1.xml",
            '<p:sld xmlns:p="p" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>Parallel speedup</a:t></p:sld>',
        )
    sources = SourceRepository(database_path); fragments = SourceFragmentRepository(database_path)

    SourceService(sources, fragments, MarkdownExtractor()).scan_source_root(vault)

    source = sources.get_by_path(source_path.resolve())
    assert source is not None and source.source_type is SourceType.PPTX
    assert [(item.text, item.location) for item in fragments.list_for_source(source.id or 0)] == [("Parallel speedup", "slide 1")]


def test_source_service_extracts_xlsx_rows_with_sheet_provenance(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "sheets"; vault.mkdir(); source_path = vault / "deadlines.xlsx"
    with ZipFile(source_path, "w") as archive:
        archive.writestr("xl/sharedStrings.xml", '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>CS3210</t></si></sst>')
        archive.writestr("xl/workbook.xml", '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Deadlines" r:id="rId1"/></sheets></workbook>')
        archive.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        archive.writestr("xl/worksheets/sheet1.xml", '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="2"><c r="A2" t="s"><v>0</v></c><c r="B2"><v>27</v></c></row></sheetData></worksheet>')
    sources = SourceRepository(database_path); fragments = SourceFragmentRepository(database_path)

    SourceService(sources, fragments, MarkdownExtractor()).scan_source_root(vault)

    source = sources.get_by_path(source_path.resolve())
    assert source is not None and source.source_type is SourceType.XLSX
    assert [(item.heading, item.text, item.location) for item in fragments.list_for_source(source.id or 0)] == [("Deadlines", "CS3210 | 27", "Deadlines!row 2")]


def test_source_service_extracts_notebook_cells_with_provenance(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "notebooks"; vault.mkdir(); source_path = vault / "queueing.ipynb"
    source_path.write_text('{"cells":[{"cell_type":"markdown","source":["# Queueing\\n"]},{"cell_type":"code","source":"lambda_rate = 2"}]}', encoding="utf-8")
    sources = SourceRepository(database_path); fragments = SourceFragmentRepository(database_path)

    SourceService(sources, fragments, MarkdownExtractor()).scan_source_root(vault)

    source = sources.get_by_path(source_path.resolve())
    assert source is not None and source.source_type is SourceType.NOTEBOOK
    assert [(item.heading, item.text, item.location) for item in fragments.list_for_source(source.id or 0)] == [("markdown", "# Queueing", "cell 1"), ("code", "lambda_rate = 2", "cell 2")]


def test_source_code_extraction_uses_stable_bounded_line_ranges(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "project"; vault.mkdir()
    code = vault / "large.py"; code.write_text("\n".join(f"line_{index} = {index}" for index in range(241)), encoding="utf-8")
    sources = SourceRepository(database_path); fragments = SourceFragmentRepository(database_path)

    SourceService(sources, fragments, MarkdownExtractor()).scan_source_root(vault)

    source = sources.get_by_path(code.resolve())
    extracted = fragments.list_for_source(source.id or 0)
    assert [item.location for item in extracted] == ["lines 1-120", "lines 121-240", "lines 241-241"]


def test_source_service_scans_and_extracts_docx_files(tmp_path: Path) -> None:
    from docx import Document

    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "vault"; vault.mkdir()
    source_path = vault / "notes.docx"
    document = Document(); document.add_heading("TLB", level=1); document.add_paragraph("Caches translations."); document.save(source_path)
    sources = SourceRepository(database_path)
    fragments = SourceFragmentRepository(database_path)
    service = SourceService(sources, fragments, MarkdownExtractor())

    result = service.scan_source_root(vault)
    source = sources.get_by_path(source_path.resolve())

    assert result.new == 1
    assert source is not None
    assert [(fragment.heading, fragment.text) for fragment in fragments.list_for_source(source.id or 0)] == [
        ("TLB", "Caches translations.")
    ]


def test_source_service_keeps_scanning_when_a_document_cannot_be_extracted(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "vault"; vault.mkdir()
    (vault / "broken.docx").write_bytes(b"not a Word document")
    (vault / "readable.txt").write_text("Still indexed", encoding="utf-8")
    sources = SourceRepository(database_path)
    fragments = SourceFragmentRepository(database_path)
    service = SourceService(sources, fragments, MarkdownExtractor())

    result = service.scan_source_root(vault)

    broken = sources.get_by_path((vault / "broken.docx").resolve())
    readable = sources.get_by_path((vault / "readable.txt").resolve())
    assert result.new == 2
    assert broken is not None
    assert fragments.list_for_source(broken.id or 0) == ()
    assert readable is not None
    assert [fragment.text for fragment in fragments.list_for_source(readable.id or 0)] == ["Still indexed"]


def test_source_service_scans_and_extracts_html_files(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "vault"; vault.mkdir()
    source_path = vault / "notes.html"; source_path.write_text("<h1>TLB</h1><p>Caches translations.</p>", encoding="utf-8")
    sources = SourceRepository(database_path)
    fragments = SourceFragmentRepository(database_path)
    service = SourceService(sources, fragments, MarkdownExtractor())

    result = service.scan_source_root(vault)
    source = sources.get_by_path(source_path.resolve())

    assert result.new == 1
    assert source is not None
    assert [(fragment.heading, fragment.text) for fragment in fragments.list_for_source(source.id or 0)] == [
        ("TLB", "Caches translations.")
    ]


def test_source_service_does_not_repeat_unchanged_image_ocr(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "vault"; vault.mkdir()
    (vault / "receipt.png").write_bytes(b"image")
    sources = SourceRepository(database_path)
    fragments = SourceFragmentRepository(database_path)
    service = SourceService(sources, fragments, MarkdownExtractor())

    class CountingOcr:
        def __init__(self) -> None:
            self.calls = 0

        def extract(self, source):
            self.calls += 1
            return ExtractionResult(
                source.id or 0,
                (SourceFragment(None, source.id or 0, None, 0, "Booking ABC123", "image OCR"),),
            )

    extractor = CountingOcr()
    service._document_extraction._extractors[SourceType.IMAGE] = extractor

    service.scan_source_root(vault)
    service.scan_source_root(vault)

    assert extractor.calls == 1
    source = sources.get_by_path((vault / "receipt.png").resolve())
    assert source is not None
    assert [fragment.text for fragment in fragments.list_for_source(source.id or 0)] == ["Booking ABC123"]


def test_source_service_reextracts_an_unchanged_source_only_when_explicitly_requested(tmp_path: Path) -> None:
    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "vault"; vault.mkdir()
    source_path = vault / "note.txt"; source_path.write_text("first", encoding="utf-8")
    sources = SourceRepository(database_path)
    fragments = SourceFragmentRepository(database_path)
    service = SourceService(sources, fragments, MarkdownExtractor())
    service.scan_source_root(vault)
    source = sources.get_by_path(source_path.resolve())
    assert source is not None

    service.reextract_source(source.id or 0)

    assert [fragment.text for fragment in fragments.list_for_source(source.id or 0)] == ["first"]


def test_source_service_keeps_scanning_when_extracted_text_cannot_be_stored(tmp_path: Path, monkeypatch) -> None:
    from steward.extraction import document

    database_path = tmp_path / "steward.db"; initialize_database(database_path)
    vault = tmp_path / "vault"; vault.mkdir()
    (vault / "a-bad.txt").write_text("bad", encoding="utf-8")
    (vault / "b-good.txt").write_text("Still indexed", encoding="utf-8")
    original = document.PlainTextExtractor.extract

    def extract(self, source):
        if source.path.name == "a-bad.txt":
            raise UnicodeEncodeError("utf-8", "\ud835", 0, 1, "surrogates not allowed")
        return original(self, source)

    monkeypatch.setattr(document.PlainTextExtractor, "extract", extract)
    sources = SourceRepository(database_path)
    fragments = SourceFragmentRepository(database_path)

    result = SourceService(sources, fragments, MarkdownExtractor()).scan_source_root(vault)

    good = sources.get_by_path((vault / "b-good.txt").resolve())
    assert result.new == 2
    assert good is not None
    assert [fragment.text for fragment in fragments.list_for_source(good.id or 0)] == ["Still indexed"]
